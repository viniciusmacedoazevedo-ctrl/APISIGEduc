import json
import re

import pytest
from fastapi.testclient import TestClient

from apisigeduc import composicoes as comp
from apisigeduc import envios
from apisigeduc import turmas as turmas_repo
from apisigeduc.modelos import Estudante, TurmaPortal
from apisigeduc.planilhas import converter, ler_conteudo, nota_de_linha
from apisigeduc.web.app import criar_app

ALUNOS = [Estudante("202630270867", "IVAN GLEYSON ZACARIAS DE OLIVEIRA"), Estudante("202630441436", "WESLEY VARELA PEREIRA DA SILVA")]


@pytest.fixture
def cenario(repo, sessoes):
    prof = repo.criar("prof@escola.br", "Prof", "senha-prof")
    turmas_repo.salvar(
        sessoes,
        prof.id,
        [
            (TurmaPortal("EFV6A", "ESCOLA HENRIQUE", "EF", "6º ANO", "História", 2), ALUNOS),
            (TurmaPortal("EFM6B", "ESCOLA DJALMA", "EF", "6º ANO", "História", 1), [Estudante("202630115378", "MARIA VITORIA")]),
        ],
    )
    turmas = {t.codigo: t for t, _ in turmas_repo.listar(sessoes, prof.id)}
    return prof, turmas


A, G = comp.AvaliacaoNova, comp.GrupoNovo


def _composicao(sessoes, prof, calculo="soma", grupos=None):
    grupos = grupos or [G("Trabalho", 2, avaliacoes=[A("Trabalho", 2, 2)]), G("Teste", 3, avaliacoes=[A("Teste", 3)]), G("Prova", 5, avaliacoes=[A("Prova", 5)])]
    return comp.criar(sessoes, prof.id, "3º bimestre - 6º anos", 2026, 3, calculo, grupos)


def _media(sessoes, prof, turma, calculo, grupos, notas_por_nome):
    """Cria a composição, lança as notas de um aluno pelo nome da avaliação e devolve a linha."""
    cid = _composicao(sessoes, prof, calculo, grupos)
    comp.aplicar(sessoes, cid, [turma.id])
    estrutura = comp.obter(sessoes, cid)
    ids = {a.nome: a.id for a in estrutura.avaliacoes}
    comp.salvar_notas(sessoes, cid, turma.id, {(ids[n], "202630115378"): v for n, v in notas_por_nome.items()})
    [linha] = comp.planilha(sessoes, cid, turma.id)
    return linha, estrutura


# Os exemplos do professor, um a um
def test_prova_unica_valendo_10(sessoes, cenario):
    prof, turmas = cenario
    linha, e = _media(sessoes, prof, turmas["EFM6B"], "soma", [G("Prova", 10, avaliacoes=[A("Prova", 10)])], {"Prova": "7,5"})
    assert linha.media == 7.5
    assert e.formula() == ["Média = Prova (10)"] and e.avisos() == []


def test_teste_10_e_prova_10_media_e_a_soma_dividida_por_2(sessoes, cenario):
    prof, turmas = cenario
    grupos = [G("Teste", 10, avaliacoes=[A("Teste", 10)]), G("Prova", 10, avaliacoes=[A("Prova", 10)])]
    linha, e = _media(sessoes, prof, turmas["EFM6B"], "media", grupos, {"Teste": "6", "Prova": "9"})
    assert linha.media == 7.5
    assert e.formula() == ["Média = média de Teste (10), Prova (10)"]


def test_teste_5_e_prova_5_media_e_a_soma(sessoes, cenario):
    prof, turmas = cenario
    grupos = [G("Teste", 5, avaliacoes=[A("Teste", 5)]), G("Prova", 5, avaliacoes=[A("Prova", 5)])]
    linha, _ = _media(sessoes, prof, turmas["EFM6B"], "soma", grupos, {"Teste": "4", "Prova": "3,5"})
    assert linha.media == 7.5


def test_trabalho_4_teste_2_prova_4(sessoes, cenario):
    prof, turmas = cenario
    grupos = [G("Trabalho", 4, avaliacoes=[A("Trabalho", 4)]), G("Teste", 2, avaliacoes=[A("Teste", 2)]), G("Prova", 4, avaliacoes=[A("Prova", 4)])]
    linha, _ = _media(sessoes, prof, turmas["EFM6B"], "soma", grupos, {"Trabalho": "4", "Teste": "1,5", "Prova": "3"})
    assert linha.media == 8.5


def test_trabalho_com_varios_trabalhos(sessoes, cenario):
    prof, turmas = cenario
    grupos = [
        G("Trabalho", 4, "soma", [A("Trabalho 1", 1), A("Trabalho 2", 1), A("Trabalho 3", 2)]),
        G("Teste", 2, avaliacoes=[A("Teste", 2)]),
        G("Prova", 4, avaliacoes=[A("Prova", 4)]),
    ]
    linha, e = _media(
        sessoes, prof, turmas["EFM6B"], "soma", grupos,
        {"Trabalho 1": "1", "Trabalho 2": "0,5", "Trabalho 3": "2", "Teste": "2", "Prova": "3"},
    )
    trabalho = e.grupos[0].grupo.id
    assert linha.grupos[trabalho] == 3.5 and linha.media == 8.5
    assert e.formula() == ["Média = Trabalho (4) + Teste (2) + Prova (4)", "Trabalho = Trabalho 1 (1) + Trabalho 2 (1) + Trabalho 3 (2)"]


def test_teste_e_prova_com_av1_e_av2(sessoes, cenario):
    prof, turmas = cenario
    grupos = [
        G("Teste", 10, "media", [A("Teste - Av1", 10), A("Teste - Av2", 10)]),
        G("Prova", 10, "media", [A("Prova - Av1", 10), A("Prova - Av2", 10)]),
    ]
    linha, e = _media(
        sessoes, prof, turmas["EFM6B"], "media", grupos,
        {"Teste - Av1": "6", "Teste - Av2": "8", "Prova - Av1": "9", "Prova - Av2": "10"},
    )
    assert [linha.grupos[g.grupo.id] for g in e.grupos] == [7.0, 9.5]
    assert linha.media == 8.3  # (7 + 9,5) / 2 = 8,25


def test_falta_uma_avaliacao_sem_media(sessoes, cenario):
    prof, turmas = cenario
    grupos = [G("Trabalho", 4, "soma", [A("Trabalho 1", 2), A("Trabalho 2", 2)]), G("Prova", 6, avaliacoes=[A("Prova", 6)])]
    linha, _ = _media(sessoes, prof, turmas["EFM6B"], "soma", grupos, {"Trabalho 1": "2", "Prova": "5"})
    assert linha.media is None


def test_validacoes():
    with pytest.raises(ValueError, match="não podem passar de 10"):
        comp.validar_estrutura("soma", [G("A", 6, avaliacoes=[A("A", 6)]), G("B", 5, avaliacoes=[A("B", 5)])])
    with pytest.raises(ValueError, match="mais que os 4 do grupo"):
        G("Trabalho", 4, "soma", [A("T1", 3), A("T2", 3)]).validar()
    with pytest.raises(ValueError, match="nota para todos"):
        G("Trabalho", 2, avaliacoes=[A("Trabalho", 2, 3)]).validar()
    with pytest.raises(ValueError, match="pelo menos uma avaliação"):
        G("Prova", 10).validar()
    assert comp.arredondar(7.25) == 7.3 and comp.arredondar(7.24) == 7.2


def test_incluir_trabalho_3_no_meio_do_bimestre(sessoes, cenario):
    prof, turmas = cenario
    grupos = [G("Trabalho", 4, "soma", [A("Trabalho 1", 2)]), G("Prova", 6, avaliacoes=[A("Prova", 6)])]
    cid = _composicao(sessoes, prof, "soma", grupos)
    e = comp.obter(sessoes, cid)
    assert e.avisos() == ["Trabalho: as avaliações somam 2 de 4 — faltam 2 ponto(s) (inclua outra avaliação quando planejar)."]
    comp.aplicar(sessoes, cid, [turmas["EFM6B"].id])
    trabalho = e.grupos[0].grupo.id
    comp.adicionar_avaliacao(sessoes, cid, trabalho, A("Trabalho 2", 2, 1.5))  # nota para todos
    e = comp.obter(sessoes, cid)
    assert [a.nome for a in e.avaliacoes] == ["Trabalho 1", "Trabalho 2", "Prova"] and e.avisos() == []
    [linha] = comp.planilha(sessoes, cid, turmas["EFM6B"].id)
    assert linha.notas == {e.avaliacoes[1].id: 1.5}
    with pytest.raises(ValueError, match="mais que os 4"):
        comp.adicionar_avaliacao(sessoes, cid, trabalho, A("Trabalho 3", 1))
    comp.remover_avaliacao(sessoes, cid, e.avaliacoes[1].id)
    assert [a.nome for a in comp.obter(sessoes, cid).avaliacoes] == ["Trabalho 1", "Prova"]


def test_nota_para_todos_e_ajuste_individual(sessoes, cenario):
    prof, turmas = cenario
    cid = _composicao(sessoes, prof)
    comp.aplicar(sessoes, cid, [turmas["EFV6A"].id, turmas["EFM6B"].id])
    trabalho, teste, prova = comp.obter(sessoes, cid).avaliacoes

    linhas = comp.planilha(sessoes, cid, turmas["EFV6A"].id)
    assert all(l.notas == {trabalho.id: 2} and l.media is None for l in linhas)  # trabalho para todos

    erros = comp.salvar_notas(
        sessoes,
        cid,
        turmas["EFV6A"].id,
        {
            (teste.id, "202630270867"): "2,5",
            (prova.id, "202630270867"): "4,25",
            (teste.id, "202630441436"): "3",
            (prova.id, "202630441436"): "7",  # passa do valor
            (trabalho.id, "202630441436"): "1",
        },
    )
    assert erros == ["WESLEY VARELA PEREIRA DA SILVA — Prova: deve ficar entre 0 e 5"]
    ivan, wesley = sorted(comp.planilha(sessoes, cid, turmas["EFV6A"].id), key=lambda l: l.nome)
    assert ivan.media == 8.8  # 2 + 2,5 + 4,25 = 8,75
    assert wesley.notas[trabalho.id] == 1 and wesley.media is None

    comp.aplicar(sessoes, cid, [turmas["EFV6A"].id])  # reaplicar não sobrescreve ajustes
    assert sorted(comp.planilha(sessoes, cid, turmas["EFV6A"].id), key=lambda l: l.nome)[1].notas[trabalho.id] == 1


def test_duplicar_para_outro_bimestre(sessoes, cenario):
    prof, _ = cenario
    grupos = [G("Teste", 10, "media", [A("Av1", 10), A("Av2", 10)]), G("Prova", 10, avaliacoes=[A("Prova", 10)])]
    cid = _composicao(sessoes, prof, "media", grupos)
    copia = comp.obter(sessoes, comp.duplicar(sessoes, cid, "4º bimestre", 4))
    assert copia.composicao.bimestre == 4 and copia.formula() == comp.obter(sessoes, cid).formula()


def test_notas_sobrevivem_a_atualizacao_da_turma(sessoes, cenario):
    prof, turmas = cenario
    cid = _composicao(sessoes, prof)
    comp.aplicar(sessoes, cid, [turmas["EFV6A"].id])
    turmas_repo.salvar(sessoes, prof.id, [(TurmaPortal("EFV6A", "ESCOLA HENRIQUE", "EF", "6º ANO", "História", 2), ALUNOS)])
    assert all(l.notas for l in comp.planilha(sessoes, cid, turmas["EFV6A"].id))


def test_csv_de_medias_e_aceito_pela_fila(sessoes, cenario):
    prof, turmas = cenario
    cid = _composicao(sessoes, prof)
    turma = turmas["EFV6A"]
    comp.aplicar(sessoes, cid, [turma.id])
    estrutura = comp.obter(sessoes, cid)
    _, teste, prova = estrutura.avaliacoes
    comp.salvar_notas(sessoes, cid, turma.id, {(teste.id, "202630270867"): "3", (prova.id, "202630270867"): "4,5"})
    dados = comp.csv_medias(estrutura.composicao, turma, comp.planilha(sessoes, cid, turma.id))
    leitura = converter(ler_conteudo("medias.csv", dados), nota_de_linha)
    assert leitura.erros == []
    [nota] = leitura.registros
    assert (nota.matricula, nota.turma, nota.disciplina, nota.bimestre, nota.nota) == ("202630270867", "EFV6A", "História", 3, 9.5)


def test_migracao_de_composicao_antiga(tmp_path):
    """Bancos da versão anterior (itens sem grupo) passam a ter um grupo por item."""
    from sqlalchemy import create_engine, text

    from apisigeduc.banco import criar_sessoes

    url = f"sqlite:///{tmp_path / 'antigo.db'}"
    engine = create_engine(url)
    with engine.begin() as c:
        c.execute(text("CREATE TABLE itens_composicao (id INTEGER PRIMARY KEY, composicao_id INTEGER, nome VARCHAR(100), valor_maximo FLOAT, nota_padrao FLOAT, ordem INTEGER)"))
        c.execute(text("INSERT INTO itens_composicao VALUES (1, 1, 'Trabalho', 2, 2, 0), (2, 1, 'Prova', 8, NULL, 1)"))
    engine.dispose()
    sessoes = criar_sessoes(url)
    with sessoes() as s:
        s.execute(text("INSERT INTO composicoes (id, usuario_id, nome, ano_letivo, bimestre, calculo) VALUES (1, 1, 'Antiga', 2026, 2, 'soma')"))
        s.commit()
    e = comp.obter(sessoes, 1)
    assert e.formula() == ["Média = Trabalho (2) + Prova (8)"]
    assert e.media({1: 2, 2: 6}) == 8.0


def _csrf(cliente, caminho):
    return re.search(r'name="csrf_token" value="([^"]+)"', cliente.get(caminho).text).group(1)


def test_fluxo_web(config, sessoes, repo, cenario):
    prof, turmas = cenario
    app = criar_app(config, sessoes, repo, iniciar_worker=False)
    cliente = TestClient(app)
    cliente.post("/login", data={"email": "prof@escola.br", "senha": "senha-prof", "csrf_token": _csrf(cliente, "/login")})
    csrf = _csrf(cliente, "/notas")

    estrutura = json.dumps([
        {"nome": "Trabalho", "valor": "2", "calculo": "soma", "avaliacoes": [{"nome": "Trabalho", "valor": "2", "padrao": "2"}]},
        {"nome": "Teste", "valor": "3", "calculo": "soma", "avaliacoes": [{"nome": "Teste", "valor": "3", "padrao": ""}]},
        {"nome": "Prova", "valor": "5", "calculo": "soma", "avaliacoes": [{"nome": "Prova", "valor": "5", "padrao": ""}]},
    ])
    r = cliente.post(
        "/notas",
        data={"nome": "3º bimestre - 6º anos", "ano_letivo": "2026", "bimestre": "3", "calculo": "soma", "csrf_token": csrf, "estrutura": estrutura},
    )
    assert "Composição criada" in r.text and "EFV6A" in r.text
    cid = int(re.search(r"/notas/(\d+)/turmas", r.text).group(1))

    r = cliente.post(f"/notas/{cid}/turmas", data={"turma_id": [turmas["EFV6A"].id, turmas["EFM6B"].id], "csrf_token": csrf})
    assert "aplicada em 2 turma(s)" in r.text

    tid = turmas["EFV6A"].id
    r = cliente.get(f"/notas/{cid}/turmas/{tid}")
    assert "IVAN GLEYSON" in r.text and 'value="2"' in r.text  # trabalho para todos
    trabalho, teste, prova = comp.obter(sessoes, cid).avaliacoes
    r = cliente.post(
        f"/notas/{cid}/turmas/{tid}",
        data={"csrf_token": csrf, f"n_{teste.id}_202630270867": "3", f"n_{prova.id}_202630270867": "4,5"},
    )
    assert "Notas salvas" in r.text and "9,5" in r.text

    # liberar para a próxima carga (não envia direto)
    r = cliente.post(f"/notas/{cid}/turmas/{tid}/liberar", data={"csrf_token": csrf})
    assert "1 estudante(s) ainda sem todas as notas" in r.text
    r = cliente.post(f"/notas/{cid}/turmas/{tid}/liberar", data={"csrf_token": csrf, "so_completos": "true"})
    assert "1 média(s) de EFV6A liberadas" in r.text and "Médias liberadas em" in r.text
    assert envios.listar(sessoes) == []  # nada vai à fila até a carga
    r = cliente.post(f"/notas/{cid}/turmas/{tid}/recolher", data={"csrf_token": csrf})
    assert "Liberação cancelada" in r.text

    # outro professor não acessa
    repo.criar("outro@escola.br", "Outro", "senha-outro")
    outro = TestClient(app)
    outro.post("/login", data={"email": "outro@escola.br", "senha": "senha-outro", "csrf_token": _csrf(outro, "/login")})
    assert outro.get(f"/notas/{cid}").status_code == 404
    assert outro.get(f"/notas/{cid}/turmas/{tid}").status_code == 404
    assert outro.post(f"/notas/{cid}/avaliacoes", data={"grupo_id": "1", "nome": "x", "valor": "1", "csrf_token": _csrf(outro, "/notas")}).status_code == 404

    # incluir avaliação pela tela: Prova passa a ter duas partes? não cabe (soma 5 de 5)
    r = cliente.post(f"/notas/{cid}/avaliacoes", data={"grupo_id": str(comp.obter(sessoes, cid).grupos[2].grupo.id), "nome": "Prova 2", "valor": "1", "csrf_token": csrf})
    assert "mais que os 5 do grupo" in r.text


def test_planilha_excel_ida_e_volta(config, sessoes, repo, cenario):
    from openpyxl import load_workbook
    import io

    from apisigeduc import planilha_notas

    prof, turmas = cenario
    grupos = [G("Trabalho", 4, "soma", [A("Trabalho 1", 2), A("Trabalho 2", 2)]), G("Prova", 6, avaliacoes=[A("Prova", 6)])]
    cid = _composicao(sessoes, prof, "soma", grupos)
    turma = turmas["EFV6A"]
    comp.aplicar(sessoes, cid, [turma.id, turmas["EFM6B"].id])
    t1, t2, prova = comp.obter(sessoes, cid).avaliacoes

    app = criar_app(config, sessoes, repo, iniciar_worker=False)
    cliente = TestClient(app)
    cliente.post("/login", data={"email": "prof@escola.br", "senha": "senha-prof", "csrf_token": _csrf(cliente, "/login")})
    r = cliente.get(f"/notas/{cid}/turmas/{turma.id}/planilha.xlsx")
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")

    # professor preenche no Excel: notas do Ivan, Wesley só o Trabalho 1, e uma nota acima do valor
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb[planilha_notas.ABA_NOTAS]
    assert ws.protection.sheet and ws.cell(planilha_notas.PRIMEIRA_LINHA, 3).protection.locked is False
    linhas = {str(ws.cell(rr, 1).value): rr for rr in range(planilha_notas.PRIMEIRA_LINHA, ws.max_row + 1)}
    ivan, wesley = linhas["202630270867"], linhas["202630441436"]
    ws.cell(ivan, 3, 2); ws.cell(ivan, 4, 1.5); ws.cell(ivan, 6, 5.5)
    ws.cell(wesley, 3, "1,5"); ws.cell(wesley, 6, 7)
    assert ws.cell(ivan, 7).value.startswith("=IF(COUNT(C")  # média é fórmula
    saida = io.BytesIO(); wb.save(saida)

    csrf = _csrf(cliente, f"/notas/{cid}/turmas/{turma.id}")
    r = cliente.post(f"/notas/{cid}/turmas/{turma.id}/importar", data={"csrf_token": csrf}, files={"arquivo": ("notas.xlsx", saida.getvalue())})
    assert "WESLEY VARELA PEREIRA DA SILVA — Prova: deve ficar entre 0 e 6" in r.text
    por_aluno = {l.matricula: l for l in comp.planilha(sessoes, cid, turma.id)}
    assert por_aluno["202630270867"].media == 9.0
    assert por_aluno["202630441436"].notas == {t1.id: 1.5}

    # planilha de outra turma é recusada
    outra = cliente.get(f"/notas/{cid}/turmas/{turmas['EFM6B'].id}/planilha.xlsx").content
    r = cliente.post(f"/notas/{cid}/turmas/{turma.id}/importar", data={"csrf_token": csrf}, files={"arquivo": ("x.xlsx", outra)})
    assert "é de outra turma ou composição (turma EFM6B)" in r.text
    r = cliente.post(f"/notas/{cid}/turmas/{turma.id}/importar", data={"csrf_token": csrf}, files={"arquivo": ("x.xlsx", b"nao e excel")})
    assert "não é uma planilha Excel" in r.text
