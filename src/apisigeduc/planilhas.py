"""Leitura de planilhas (.csv ou .xlsx) e conversão para os modelos."""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Generic, Iterable, TypeVar

from .modelos import Frequencia, Nota

T = TypeVar("T")


@dataclass
class ResultadoLeitura(Generic[T]):
    registros: list[T] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)


def _normalizar_cabecalho(nome: str) -> str:
    return nome.strip().lower().replace(" ", "_")


def ler_linhas(caminho: Path) -> list[dict[str, str]]:
    """Lê um .csv (separado por ; ou ,) ou .xlsx e devolve uma lista de dicts."""
    sufixo = caminho.suffix.lower()
    if sufixo == ".csv":
        with caminho.open(encoding="utf-8-sig", newline="") as f:
            amostra = f.read(4096)
            f.seek(0)
            dialeto = csv.Sniffer().sniff(amostra, delimiters=";,")
            leitor = csv.DictReader(f, dialect=dialeto)
            return [
                {_normalizar_cabecalho(k): (v or "").strip() for k, v in linha.items() if k}
                for linha in leitor
            ]
    if sufixo in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        planilha = load_workbook(caminho, read_only=True, data_only=True).active
        linhas = planilha.iter_rows(values_only=True)
        cabecalho = [_normalizar_cabecalho(str(c or "")) for c in next(linhas)]
        return [
            {k: "" if v is None else str(v).strip() for k, v in zip(cabecalho, linha) if k}
            for linha in linhas
            if any(v not in (None, "") for v in linha)
        ]
    raise ValueError(f"Formato não suportado: {caminho.suffix} (use .csv ou .xlsx)")


def _numero(texto: str) -> float:
    return float(texto.replace(",", "."))


def _data(texto: str) -> date:
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError(f"data inválida: {texto!r} (use DD/MM/AAAA)")


def _presenca(texto: str) -> bool:
    valor = texto.strip().upper()
    if valor in {"P", "PRESENTE", "1", "SIM", "S", "TRUE"}:
        return True
    if valor in {"F", "FALTA", "AUSENTE", "0", "NAO", "NÃO", "N", "FALSE"}:
        return False
    raise ValueError(f"presença inválida: {texto!r} (use P ou F)")


def nota_de_linha(linha: dict[str, str]) -> Nota:
    return Nota(
        matricula=linha["matricula"],
        turma=linha["turma"],
        disciplina=linha["disciplina"],
        bimestre=int(_numero(linha["bimestre"])),
        nota=_numero(linha["nota"]),
    )


def frequencia_de_linha(linha: dict[str, str]) -> Frequencia:
    return Frequencia(
        matricula=linha["matricula"],
        turma=linha["turma"],
        disciplina=linha["disciplina"],
        data=_data(linha["data"]),
        presente=_presenca(linha["presenca"]),
    )


def converter(
    linhas: Iterable[dict[str, str]], conversor: Callable[[dict[str, str]], T]
) -> ResultadoLeitura[T]:
    """Converte e valida cada linha, acumulando erros em vez de parar no primeiro."""
    resultado: ResultadoLeitura[T] = ResultadoLeitura()
    for numero, linha in enumerate(linhas, start=2):  # linha 1 é o cabeçalho
        try:
            registro = conversor(linha)
            registro.validar()  # type: ignore[attr-defined]
            resultado.registros.append(registro)
        except KeyError as e:
            resultado.erros.append(f"linha {numero}: coluna ausente {e}")
        except ValueError as e:
            resultado.erros.append(f"linha {numero}: {e}")
    return resultado
