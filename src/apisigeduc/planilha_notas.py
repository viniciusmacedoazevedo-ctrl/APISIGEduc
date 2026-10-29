"""Planilha Excel da grade de notas de uma turma: exportar para o professor
preencher fora do sistema e importar de volta.

A nota de cada grupo e a média são FÓRMULAS do Excel (mesma regra e mesmo
arredondamento do sistema), então a planilha recalcula enquanto o professor
digita. Só as células amarelas são editáveis. Uma aba oculta guarda a
identificação da composição e da turma, para não importar a planilha errada.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import composicoes as comp
from .banco import Turma

ABA_NOTAS = "Notas"
ABA_AJUDA = "Como preencher"
ABA_DADOS = "_apisigeduc"
LINHA_GRUPOS, LINHA_AVALIACOES, LINHA_VALE, LINHA_IDS = 4, 5, 6, 7
PRIMEIRA_LINHA = 8

_FONTE = "Arial"
_AMARELO = PatternFill("solid", start_color="FFF2CC")
_CINZA = PatternFill("solid", start_color="EEF2EF")
_FINA = Side(style="thin", color="C9D3CC")
_BORDA = Border(left=_FINA, right=_FINA, top=_FINA, bottom=_FINA)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fonte(**kw) -> Font:
    return Font(name=_FONTE, **kw)


def exportar(estrutura: comp.Estrutura, turma: Turma, linhas: list[comp.LinhaAluno]) -> bytes:
    c = estrutura.composicao
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_NOTAS

    ws["A1"] = f"{turma.codigo} {turma.componente}: {c.nome}"
    ws["A1"].font = _fonte(bold=True, size=14)
    ws["A2"] = "Preencha só as células amarelas. A nota de cada grupo e a média são calculadas pelas fórmulas."
    ws["A2"].font = _fonte(italic=True, color="56625B")
    ws["A3"] = "; ".join(estrutura.formula())
    ws["A3"].font = _fonte(color="56625B", size=9)

    for linha, rotulo in ((LINHA_AVALIACOES, "Matrícula"), (LINHA_VALE, "Vale")):
        ws.cell(linha, 1, rotulo)
    ws.cell(LINHA_AVALIACOES, 2, "Estudante")

    # colunas: avaliações de cada grupo + "nota do grupo" quando há várias
    col = 3
    colunas_grupo = []  # (grupo, colunas das avaliações, coluna da nota do grupo)
    for g in estrutura.grupos:
        cols_aval = []
        inicio = col
        for a in g.avaliacoes:
            ws.cell(LINHA_AVALIACOES, col, a.nome if len(g.avaliacoes) > 1 else g.grupo.nome)
            ws.cell(LINHA_VALE, col, a.valor_maximo)  # constante: valor máximo da avaliação
            ws.cell(LINHA_IDS, col, a.id)
            cols_aval.append(col)
            col += 1
        col_grupo = None
        if len(g.avaliacoes) > 1:
            col_grupo = col
            ws.cell(LINHA_AVALIACOES, col, f"Nota {g.grupo.nome}")
            ws.cell(LINHA_VALE, col, g.grupo.valor)
            col += 1
        ws.cell(LINHA_GRUPOS, inicio, f"{g.grupo.nome} (vale {comp._fmt(g.grupo.valor)}"
                + (f", {comp.CALCULOS[g.grupo.calculo].lower()})" if len(g.avaliacoes) > 1 else ")"))
        if col - 1 > inicio:
            ws.merge_cells(start_row=LINHA_GRUPOS, start_column=inicio, end_row=LINHA_GRUPOS, end_column=col - 1)
        colunas_grupo.append((g, cols_aval, col_grupo))
    col_media = col
    ws.cell(LINHA_AVALIACOES, col_media, "Média")
    ultima_coluna = col_media

    # estilo do cabeçalho
    for linha in (LINHA_GRUPOS, LINHA_AVALIACOES, LINHA_VALE):
        for coluna in range(1, ultima_coluna + 1):
            celula = ws.cell(linha, coluna)
            fixo = linha == LINHA_VALE and coluna > 2 and celula.value is not None
            celula.font = _fonte(bold=linha != LINHA_VALE, color="0000FF" if fixo else "17211C")
            celula.fill = _CINZA
            celula.alignment = _CENTRO
            celula.border = _BORDA
    ws.row_dimensions[LINHA_IDS].hidden = True

    todas = [cc for _, cols, _ in colunas_grupo for cc in cols]
    n = len(todas)
    for i, aluno in enumerate(linhas):
        r = PRIMEIRA_LINHA + i
        m = ws.cell(r, 1, aluno.matricula)
        m.number_format = "@"
        ws.cell(r, 2, aluno.nome)
        for coluna in (1, 2):
            ws.cell(r, coluna).font = _fonte()
            ws.cell(r, coluna).border = _BORDA
        partes_media = []  # (referência da nota do grupo, referência do valor do grupo)
        for g, cols, col_grupo in colunas_grupo:
            for coluna, a in zip(cols, g.avaliacoes):
                celula = ws.cell(r, coluna, aluno.notas.get(a.id))
                celula.fill = _AMARELO
                celula.protection = Protection(locked=False)
                celula.number_format = "0.0#"
                celula.font = _fonte(color="0000FF")
                celula.alignment = _CENTRO
                celula.border = _BORDA
            refs = [f"{get_column_letter(cc)}{r}" for cc in cols]
            if col_grupo is None:
                partes_media.append((refs[0], f"{get_column_letter(cols[0])}${LINHA_VALE}"))
                continue
            faixa = f"{refs[0]}:{refs[-1]}"
            if g.grupo.calculo == "media":
                termos = "+".join(f"{ref}/{get_column_letter(cc)}${LINHA_VALE}" for ref, cc in zip(refs, cols))
                expr = f"({termos})/{len(refs)}*{get_column_letter(col_grupo)}${LINHA_VALE}"
            else:
                expr = f"SUM({faixa})"
            celula = ws.cell(r, col_grupo, f'=IF(COUNT({faixa})<{len(refs)},"",{expr})')
            celula.number_format = "0.0#"
            celula.font = _fonte(color="56625B")
            celula.alignment = _CENTRO
            celula.border = _BORDA
            partes_media.append((f"{get_column_letter(col_grupo)}{r}", f"{get_column_letter(col_grupo)}${LINHA_VALE}"))

        todas_refs = ",".join(f"{get_column_letter(cc)}{r}" for cc in todas)
        if c.calculo == "media":
            expr = "(" + "+".join(f"{nota}/{vale}" for nota, vale in partes_media) + f")/{len(partes_media)}*10"
        else:
            expr = "+".join(nota for nota, _ in partes_media)
        # ROUND do Excel arredonda 0,25 para 0,3, como o sistema
        celula = ws.cell(r, col_media, f'=IF(COUNT({todas_refs})<{n},"",ROUND(MIN({expr},10),1))')
        celula.number_format = "0.0"
        celula.font = _fonte(bold=True)
        celula.alignment = _CENTRO
        celula.border = _BORDA

    # validação: nota entre 0 e o valor da avaliação
    if linhas:
        ultima = PRIMEIRA_LINHA + len(linhas) - 1
        for g, cols, _ in colunas_grupo:
            for coluna, a in zip(cols, g.avaliacoes):
                dv = DataValidation(
                    type="decimal", operator="between", formula1="0", formula2=str(a.valor_maximo), allow_blank=True,
                    showErrorMessage=True, errorTitle="Nota inválida",
                    error=f"{a.nome}: digite uma nota de 0 a {comp._fmt(a.valor_maximo)}.",
                )
                ws.add_data_validation(dv)
                letra = get_column_letter(coluna)
                dv.add(f"{letra}{PRIMEIRA_LINHA}:{letra}{ultima}")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40
    for coluna in range(3, ultima_coluna + 1):
        ws.column_dimensions[get_column_letter(coluna)].width = 13
    ws.row_dimensions[LINHA_GRUPOS].height = 30
    ws.row_dimensions[LINHA_AVALIACOES].height = 30
    ws.freeze_panes = ws.cell(PRIMEIRA_LINHA, 3)
    ws.protection.sheet = True  # só as células amarelas podem ser editadas
    ws.protection.formatColumns = False

    _aba_ajuda(wb, estrutura)
    dados = wb.create_sheet(ABA_DADOS)
    for i, (chave, valor) in enumerate(
        (("composicao_id", c.id), ("turma_id", turma.id), ("turma", turma.codigo), ("versao", 1)), start=1
    ):
        dados.cell(i, 1, chave)
        dados.cell(i, 2, valor)
    dados.sheet_state = "hidden"

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _aba_ajuda(wb: Workbook, estrutura: comp.Estrutura) -> None:
    ws = wb.create_sheet(ABA_AJUDA)
    linhas = [
        ("Como preencher esta planilha", True),
        ("", False),
        ("1. Na aba Notas, digite as notas só nas células amarelas (uma coluna por avaliação).", False),
        ("2. Cada nota vai de 0 até o valor da linha Vale. O Excel recusa valores fora disso.", False),
        ("3. A nota de cada grupo e a média são calculadas sozinhas: não precisa digitar.", False),
        ("4. Pode deixar células em branco: na importação, elas não apagam as notas já lançadas.", False),
        ("5. Salve e, no sistema, abra a turma em Notas e use Importar planilha.", False),
        ("", False),
        ("Como a média é calculada", True),
        *[(linha, False) for linha in estrutura.formula()],
        ("A média é arredondada para 1 casa decimal (8,25 vira 8,3).", False),
        ("", False),
        ("Exemplo de preenchimento (não faz parte da turma)", True),
    ]
    for i, (texto, negrito) in enumerate(linhas, start=1):
        ws.cell(i, 1, texto).font = _fonte(bold=negrito, size=12 if negrito else 11)
    exemplo = len(linhas) + 1
    cabecalho = ["Matrícula", "Estudante"] + [a.nome for a in estrutura.avaliacoes]
    valores = ["202600000000", "NOME DO ESTUDANTE"] + [comp._fmt(round(a.valor_maximo * 0.8, 1)) for a in estrutura.avaliacoes]
    for j, (cab, val) in enumerate(zip(cabecalho, valores), start=1):
        ws.cell(exemplo, j, cab).font = _fonte(bold=True)
        ws.cell(exemplo, j).fill = _CINZA
        celula = ws.cell(exemplo + 1, j, val)
        celula.font = _fonte(color="0000FF" if j > 2 else "17211C")
        if j > 2:
            celula.fill = _AMARELO
    ws.cell(exemplo + 3, 1, "Cores: amarelo = você preenche; texto azul = valores digitados ou fixos; preto = calculado.").font = _fonte(italic=True, color="56625B")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 30


class PlanilhaInvalida(ValueError):
    pass


def importar(estrutura: comp.Estrutura, turma: Turma, dados: bytes) -> dict[tuple[int, str], str]:
    """Lê a planilha preenchida. Devolve {(avaliação, matrícula): nota} só com as
    células preenchidas (célula em branco não apaga nota)."""
    try:
        wb = load_workbook(io.BytesIO(dados), data_only=True)
    except Exception:
        raise PlanilhaInvalida("O arquivo não é uma planilha Excel (.xlsx) válida.") from None
    if ABA_NOTAS not in wb.sheetnames or ABA_DADOS not in wb.sheetnames:
        raise PlanilhaInvalida("Use a planilha baixada em \"Baixar planilha (Excel)\" desta turma.")
    meta = {str(r[0]): r[1] for r in wb[ABA_DADOS].iter_rows(values_only=True) if r and r[0]}
    if meta.get("composicao_id") != estrutura.composicao.id or meta.get("turma_id") != turma.id:
        raise PlanilhaInvalida(
            f"Esta planilha é de outra turma ou composição (turma {meta.get('turma', '?')}). "
            f"Baixe a planilha desta turma ({turma.codigo})."
        )
    ws = wb[ABA_NOTAS]
    validas = {a.id for a in estrutura.avaliacoes}
    colunas = {
        coluna: valor
        for coluna, valor in enumerate(next(ws.iter_rows(min_row=LINHA_IDS, max_row=LINHA_IDS, values_only=True)), start=1)
        if isinstance(valor, int) and valor in validas
    }
    valores: dict[tuple[int, str], str] = {}
    for linha in ws.iter_rows(min_row=PRIMEIRA_LINHA, values_only=True):
        matricula = linha[0]
        if matricula in (None, ""):
            continue
        matricula = str(matricula).strip()
        for coluna, item_id in colunas.items():
            valor = linha[coluna - 1] if coluna - 1 < len(linha) else None
            if valor is None or (isinstance(valor, str) and not valor.strip()):
                continue
            valores[(item_id, matricula)] = str(valor).replace(".", ",") if isinstance(valor, (int, float)) else str(valor)
    return valores
